"""
Generate the companion Excel workbook: futures_risk_modeler.xlsx
Run: python3 build_excel.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

C_DARK   = "0B1120"; C_NAVY  = "1A1A2E"; C_PANEL = "111827"
C_ROWALT = "0F172A"; C_INPUT = "0D2137"; C_RESULT = "0F2910"
C_WARN_BG= "2D1B00"
C_WHITE  = "F8FAFC"; C_SLATE = "94A3B8"
C_INPUT_T= "7DD3FC"; C_RES_T = "86EFAC"; C_WARN_T = "FCD34D"

def fill(h): return PatternFill("solid", fgColor=h)
def fnt(h, bold=False, sz=10): return Font(color=h, bold=bold, size=sz, name="Calibri")
def bdr():
    s = Side(style="thin", color="1E293B")
    return Border(left=s, right=s, top=s, bottom=s)

def cell(ws, r, c, val="", bg=C_PANEL, fg=C_WHITE, bold=False, num=None, align="left", sz=10):
    cc = ws.cell(row=r, column=c, value=val)
    cc.fill = fill(bg); cc.font = fnt(fg, bold=bold, sz=sz); cc.border = bdr()
    cc.alignment = Alignment(horizontal=align, vertical="center")
    if num: cc.number_format = num
    return cc

def sec(ws, r, c1, text, c2=None):
    """Section label. c2 = last column to merge into (not span count)."""
    cc = cell(ws, r, c1, text, bg=C_NAVY, fg=C_SLATE, bold=True, sz=9)
    ws.row_dimensions[r].height = 15
    if c2 and c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        ws.cell(r, c1).alignment = Alignment(horizontal="left", vertical="center")
    return cc

def title_row(ws, text, last_col):
    ws.row_dimensions[1].height = 28
    c = ws.cell(row=1, column=1, value=text)
    ws.merge_cells(f"A1:{get_column_letter(last_col)}1")
    c.fill = fill(C_NAVY)
    c.font = Font(color=C_WHITE, bold=True, size=14, name="Calibri")
    c.alignment = Alignment(horizontal="left", vertical="center")

# ═══════════════════════════════════════════════════════════════════
# SHEET 1 — Risk Calculator
# ═══════════════════════════════════════════════════════════════════
def risk_calc(wb):
    ws = wb.create_sheet("Risk Calculator")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0EA5E9"
    for col, w in zip("ABCDE", [26, 16, 4, 22, 16]):
        ws.column_dimensions[col].width = w

    title_row(ws, "📐  POSITION SIZER", 5)

    # LEFT: Inputs (cols A-B)
    sec(ws, 3, 1, "INPUTS", 2)
    inputs = [
        ("Account Size ($)",       5000,  "$#,##0.00"),
        ("Risk (%)",               0.5,   "0.00"),
        ("Stop Loss (ticks)",      6,     "0"),
        ("Tick Value ($/tick)",    1.25,  "$#,##0.00"),
        ("Commission / side ($)",  0.35,  "$#,##0.00"),
        ("Slippage (ticks)",       0.5,   "0.00"),
    ]
    for i, (lbl, val, fmt) in enumerate(inputs, start=4):
        ws.row_dimensions[i].height = 18
        cell(ws, i, 1, lbl, bg="1E293B", fg=C_SLATE)
        cell(ws, i, 2, val, bg=C_INPUT, fg=C_INPUT_T, bold=True, num=fmt, align="right")

    # LEFT: Results (cols A-B)
    sec(ws, 11, 1, "RESULTS  (auto-calculated)", 2)
    # B4=account, B5=risk%, B6=stop_ticks, B7=tick_val, B8=comm/side, B9=slip_ticks
    results = [
        ("Dollar Risk ($)",           "=B4*(B5/100)",                      "$#,##0.00"),
        ("Stop Cost / Contract ($)",  "=B6*B7",                            "$#,##0.00"),
        ("Friction / Contract ($)",   "=(B8*2)+(B9*B7*2)",                "$#,##0.00"),
        ("Total Cost / Contract ($)", "=B13+B14",                          "$#,##0.00"),
        ("Max Contracts",             "=IF(B15>0,INT(B12/B15),0)",         "0"),
        ("Effective Risk ($)",        "=B16*B15",                          "$#,##0.00"),
        ("Effective Risk (%)",        "=IFERROR(B17/B4*100,0)",            "0.00"),
    ]
    for i, (lbl, formula, fmt) in enumerate(results, start=12):
        ws.row_dimensions[i].height = 18
        cell(ws, i, 1, lbl, bg="1E293B", fg=C_SLATE)
        cell(ws, i, 2, formula, bg=C_RESULT, fg=C_RES_T, bold=True, num=fmt, align="right")

    ws.row_dimensions[20].height = 18
    cell(ws, 20, 1, "Check:", bg="1E293B", fg=C_SLATE)
    cell(ws, 20, 2,
         '=IF(B5>2,"⚠️ Risk>2%",IF(B16=0,"🚫 Cannot size 1 contract","✅ Size OK"))',
         bg=C_WARN_BG, fg=C_WARN_T, bold=True, align="center")

    # RIGHT: Instrument reference (cols D-E)
    sec(ws, 3, 4, "INSTRUMENT REFERENCE", 5)
    for ci, h in enumerate(["Instrument", "Tick Value"], start=4):
        ws.row_dimensions[4].height = 18
        cell(ws, 4, ci, h, bg=C_NAVY, fg=C_WHITE, bold=True, align="center")

        ref = [("MES","$1.25/tick"),("MGC","$1.00/tick"),("MNQ","$0.50/tick"),("M2K","$0.50/tick")]
    for ri, (sym, tv) in enumerate(ref, start=5):
        ws.row_dimensions[ri].height = 18
        bg = "1A2744" if ri % 2 == 0 else C_PANEL
        cell(ws, ri, 4, sym, bg=bg, fg=C_WHITE, bold=True, align="center")
        cell(ws, ri, 5, tv,  bg=bg, fg=C_RES_T, align="center")

    sec(ws, 10, 4, "FORMULAS", 5)
    tips = [
        "dollar_risk = account × (risk% / 100)",
        "stop_cost = stop_ticks × tick_value",
        "friction = 2×comm + 2×(slip_ticks×tick_val)",
        "total_cost = stop_cost + friction",
        "max_contracts = FLOOR(dollar_risk / total_cost)",
    ]
    for ri, tip in enumerate(tips, start=11):
        ws.row_dimensions[ri].height = 16
        ws.merge_cells(start_row=ri, start_column=4, end_row=ri, end_column=5)
        cc = cell(ws, ri, 4, tip, bg=C_DARK, fg=C_SLATE, sz=9)
        cc.alignment = Alignment(horizontal="left", vertical="center")


# ═══════════════════════════════════════════════════════════════════
# SHEET 2 — Prop Firm Tracker
# ═══════════════════════════════════════════════════════════════════
def prop_tracker(wb):
    ws = wb.create_sheet("Prop Firm Tracker")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "8B5CF6"
    for col, w in zip("ABCDEFG", [26, 16, 4, 12, 16, 16, 16]):
        ws.column_dimensions[col].width = w

    title_row(ws, "🏦  PROP FIRM / ACCOUNT TRACKER", 7)

    sec(ws, 3, 1, "ACCOUNT CONFIG  (edit blue cells)", 2)
    cfg = [
        ("Firm Name",         "Custom",   None),
        ("Starting Balance",  50000,      "$#,##0"),
        ("Max Drawdown",      2500,       "$#,##0"),
        ("Daily Loss Limit",  1000,       "$#,##0"),
        ("Profit Target",     3000,       "$#,##0"),
        ("Drawdown Type",     "static",   None),
        ("Current Balance",   50000,      "$#,##0"),
        ("Realized P/L",      0,          "$#,##0.00"),
        ("High Water Mark",   50000,      "$#,##0"),
    ]
    for i, (lbl, val, fmt) in enumerate(cfg, start=4):
        ws.row_dimensions[i].height = 18
        cell(ws, i, 1, lbl, bg="1E293B", fg=C_SLATE)
        cell(ws, i, 2, val, bg=C_INPUT, fg=C_INPUT_T, bold=True, num=fmt, align="right")

    # B4=firm_name B5=start_bal B6=max_dd B7=daily_limit B8=profit_tgt
    # B9=dd_type B10=curr_bal B11=realized_pnl B12=hwm
    sec(ws, 14, 1, "DERIVED  (auto-calculated)", 2)
    derived = [
        ("Floor Balance ($)",
         '=IF(LOWER(B9)="trailing",B12-B6,B5-B6)',                          "$#,##0"),
        ("Remaining Buffer ($)",
         "=B10-B15",                                                          "$#,##0"),
        ("Buffer % of Max DD",
         "=IFERROR(B16/B6,0)",                                               "0.0%"),
        ("Distance to Target ($)",
         "=MAX(0,B8-B11)",                                                    "$#,##0"),
        ("Account Status",
         '=IF(B16<=0,"🔴 FAILED",IF(B17<0.1,"🟠 DANGER",IF(B17<0.3,"🟡 CAUTION","🟢 ON TRACK")))',
         None),
    ]
    for i, (lbl, formula, fmt) in enumerate(derived, start=15):
        ws.row_dimensions[i].height = 18
        cell(ws, i, 1, lbl, bg="1E293B", fg=C_SLATE)
        cc = cell(ws, i, 2, formula, bg=C_RESULT, fg=C_RES_T, bold=True, align="right")
        if fmt: cc.number_format = fmt

    sec(ws, 22, 1, "DAILY SESSION LOG", 7)
    hdrs = ["Date", "Start $", "P/L $", "End $", "Buffer $", "Limit Used", "Notes"]
    ws.row_dimensions[23].height = 18
    for ci, h in enumerate(hdrs, start=1):
        cell(ws, 23, ci, h, bg=C_NAVY, fg=C_WHITE, bold=True, align="center")

    for ri in range(24, 74):
        ws.row_dimensions[ri].height = 17
        bg = C_ROWALT if ri % 2 == 0 else C_DARK
        cell(ws, ri, 1, "", bg=bg, fg=C_WHITE, num="YYYY-MM-DD")
        cell(ws, ri, 2, "", bg=bg, fg=C_INPUT_T, num="$#,##0.00")
        cell(ws, ri, 3, "", bg=bg, fg=C_INPUT_T, num="$#,##0.00")
        cell(ws, ri, 4, f"=IF(OR(B{ri}=0,C{ri}=0),\"\",B{ri}+C{ri})",
             bg=bg, fg=C_RES_T, num="$#,##0.00", align="right")
        cell(ws, ri, 5, f'=IF(D{ri}="","",D{ri}-$B$15)',
             bg=bg, fg=C_RES_T, num="$#,##0.00", align="right")
        cell(ws, ri, 6, f'=IF(C{ri}=0,"",ABS(C{ri})/$B$7)',
             bg=bg, fg=C_WARN_T, num="0%", align="center")
        cell(ws, ri, 7, "", bg=bg, fg=C_SLATE)


# ═══════════════════════════════════════════════════════════════════
# SHEET 3 — Trade Journal
# ═══════════════════════════════════════════════════════════════════
def trade_journal(wb):
    ws = wb.create_sheet("Trade Journal")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "10B981"
    ws.freeze_panes = "A3"

    widths = [6, 12, 20, 10, 22, 9, 9, 9, 9, 5, 9, 10, 18, 16, 7, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    title_row(ws, "📓  TRADE JOURNAL", len(widths))

    hdrs = ["ID", "Date", "Session", "Instrument", "Setup",
            "Entry", "Stop", "Target", "Exit", "Contr",
            "R", "P/L $", "Mistake", "Emotion", "Grade", "Notes"]
    ws.row_dimensions[2].height = 20
    for ci, h in enumerate(hdrs, start=1):
        cell(ws, 2, ci, h, bg=C_NAVY, fg=C_WHITE, bold=True, align="center", sz=9)

    col_fmts = [
        "0", "YYYY-MM-DD", None, None, None,
        "$#,##0.00", "$#,##0.00", "$#,##0.00", "$#,##0.00", "0",
        "0.00", "$#,##0.00", None, None, None, None,
    ]
    col_fgs = [
        C_SLATE, C_WHITE, C_WHITE, C_WHITE, C_WHITE,
        C_SLATE, C_SLATE, C_SLATE, C_SLATE, C_SLATE,
        C_INPUT_T, C_RES_T, C_WHITE, C_WHITE, C_WARN_T, C_SLATE,
    ]
    for ri in range(3, 203):
        ws.row_dimensions[ri].height = 17
        bg = C_ROWALT if ri % 2 == 0 else C_DARK
        for ci in range(1, len(hdrs) + 1):
            cc = cell(ws, ri, ci, "", bg=bg, fg=col_fgs[ci-1], sz=9)
            if col_fmts[ci-1]:
                cc.number_format = col_fmts[ci-1]

    # Data validation — FIX: pass string to dv.add(), not ws[range]
    dv_map = {
        "C": '"New York,London,Overlap (London/NY),Asian,Afternoon,Other"',
        "D": '"MES,MGC,MNQ,M2K"',
        "E": '"VWAP Bounce,VWAP Rejection,SR Break,SR Rejection,EMA Pullback,Liquidity Sweep,Opening Range Break,Trend Continuation,Mean Reversion,Other"',
        "M": '"None,FOMO,Revenge Trade,Early Entry,Late Entry,Missed Exit,Chased Price,Ignored Stop,Oversized,Broke Rules,Other"',
        "N": '"Calm,Confident,Anxious,Bored,Frustrated,FOMO,Fearful,Greedy,Distracted,Other"',
        "O": '"A,B,C,D"',
    }
    for col_letter, formula in dv_map.items():
        dv = DataValidation(type="list", formula1=formula,
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}3:{col_letter}202")   # ← FIX: string, not ws[range]


# ═══════════════════════════════════════════════════════════════════
# SHEET 4 — Instructions
# ═══════════════════════════════════════════════════════════════════
def instructions(wb):
    ws = wb.create_sheet("Instructions")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "F59E0B"
    ws.column_dimensions["A"].width = 82

    title_row(ws, "ℹ️  HOW TO USE THIS COMPANION SPREADSHEET", 1)

    lines = [
        ("", None),
        ("📐  RISK CALCULATOR", "h"),
        ("  Edit the blue input cells. All results update automatically.", "b"),
        ("  Tick values:  MES=$1.25  ·  MGC=$1.00  ·  MNQ=$0.50  ·  M2K=$0.50 per tick.", "b"),
        ("  'Tick Value' input: copy from the Instrument Reference table to the right.", "b"),
        ("  For full interactivity, run the Streamlit app:  streamlit run app.py", "b"),
        ("", None),
        ("🏦  PROP FIRM TRACKER", "h"),
        ("  1. Fill in Account Config once at challenge start (blue cells only).", "b"),
        ("  2. Update Current Balance and Realized P/L each trading day.", "b"),
        ("  3. Trailing DD: update High Water Mark any time you reach a new balance peak.", "b"),
        ("  4. Daily Session Log: enter Start $ and P/L $ — all other columns are formulas.", "b"),
        ("  5. Account Status: 🟢 ON TRACK / 🟡 CAUTION (<30% buffer) / 🟠 DANGER (<10%) / 🔴 FAILED", "b"),
        ("", None),
        ("📓  TRADE JOURNAL", "h"),
        ("  1. Log every trade. Session, Instrument, Setup, Mistake, Emotion, Grade have dropdowns.", "b"),
        ("  2. Result R: positive = win (e.g. 1.5), negative = loss (e.g. -1.0).", "b"),
        ("  3. Grade:  A = textbook setup, perfect execution", "b"),
        ("           B = good setup or minor execution issue", "b"),
        ("           C = marginal setup — should be reviewed", "b"),
        ("           D = should not have taken this trade", "b"),
        ("  4. Export to CSV → drag file into Streamlit app → full analytics dashboard.", "b"),
        ("", None),
        ("📊  ANALYTICS", "h"),
        ("  This spreadsheet handles daily tracking. For edge validation, equity curves,", "b"),
        ("  session breakdown, and R distribution, use the Streamlit app.", "b"),
        ("  Run: streamlit run app.py  →  Analytics tab", "b"),
        ("", None),
        ("💡  RULES TO INTERNALIZE", "h"),
        ("  • Risk ≤ 0.5% per trade during evaluation phase. 1% max once funded.", "b"),
        ("  • Never risk more than 25% of remaining drawdown buffer in a single day.", "b"),
        ("  • Stop trading after using 50% of your daily loss limit.", "b"),
        ("  • Log every trade, including ones you want to forget.", "b"),
        ("  • Do not size up before 30+ trades with positive expectancy.", "b"),
        ("  • If Grade C/D trades are negative expectancy, stop taking them.", "b"),
    ]

    for ri, (text, kind) in enumerate(lines, start=2):
        ws.row_dimensions[ri].height = 17
        cc = ws.cell(row=ri, column=1, value=text)
        if kind == "h":
            cc.font = Font(color="93C5FD", bold=True, size=11, name="Calibri")
            cc.fill = fill("111827")
        elif kind == "b":
            cc.font = Font(color="CBD5E1", size=10, name="Calibri")
            cc.fill = fill(C_DARK)
        else:
            cc.font = Font(color="475569", size=10, name="Calibri")
            cc.fill = fill(C_DARK)
        cc.alignment = Alignment(horizontal="left", vertical="center")


# ── Assemble ─────────────────────────────────────────────────────────────────
wb = Workbook()
wb.remove(wb.active)
risk_calc(wb)
prop_tracker(wb)
trade_journal(wb)
instructions(wb)

out = "futures_risk_modeler.xlsx"
wb.save(out)

# Quick verification
from openpyxl import load_workbook
wb2 = load_workbook(out)
print(f"Saved: {out}")
print(f"Sheets: {wb2.sheetnames}")
# Spot-check a formula
ws1 = wb2["Risk Calculator"]
print(f"Max Contracts formula: {ws1['B16'].value}")
ws2 = wb2["Prop Firm Tracker"]
print(f"Floor Balance formula: {ws2['B15'].value}")
ws3 = wb2["Trade Journal"]
print(f"Dropdowns attached: {len(ws3.data_validations.dataValidation)} validations")
print("Excel build: OK")
